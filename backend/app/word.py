"""The tome as an editable Word document: text, lists, tables and exercises stay text the teacher can change;
what the code draws (fichas, diagrams, tracing, keyboards) and the illustrations go in as pictures.

The drawings are the same SVG the PDF prints, rasterised in one WeasyPrint pass (one page per drawing, each
page the drawing's size) and cut into PNGs with pdfium, so Word shows exactly what the PDF shows.
"""

import io
import re

import pypdfium2
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

from app.content import DIAGRAM_TYPES, WORKSHEET_TYPES
from app.diagrams import render_diagram
from app.printing import PAPERS
from app.render import ASSETS, GENRE_LABEL, Tome, keyboard_svg, letter_grid_cells, theme_colors, tracing_svg
from app.themes import FONTS
from app.worksheets import answers as worksheet_answers, render_worksheet

MARGIN_MM = 18
DRAWN = set(DIAGRAM_TYPES) | set(WORKSHEET_TYPES) | {"keyboard", "tracing"}
CALLOUT_LABEL = {"key_idea": "Idea clave", "tip": "Consejo", "warning": "Atención", "fun_fact": "¿Sabías que…?"}
VIEWBOX = re.compile(r'viewBox="([-\d.]+) ([-\d.]+) ([\d.]+) ([\d.]+)"')
PNG_WIDTH = 1800  # px across the text column: about 250 dpi at A4


def render_docx(tome: Tome, *, paper: str, include_answer_key: bool) -> bytes:
    th, c = theme_colors(tome.design, tome.palette)
    accent = RGBColor.from_string(c["accent"].lstrip("#"))
    doc = Document()
    w, h = PAPERS[paper]
    for s in doc.sections:
        s.page_width, s.page_height = Mm(w), Mm(h)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Mm(MARGIN_MM)
    column = Mm(w - 2 * MARGIN_MM)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = th.body, Pt(th.body_size_pt)
    for name, size in (("Title", 32), ("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12.5)):
        style = doc.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = th.head, Pt(size), accent
    drawings = _rasterise(tome, c, th.head)

    # Cover: the painted art, then the type the PDF sets over it.
    if tome.cover_image:
        doc.add_picture(io.BytesIO(_word_image(tome.cover_image)), width=column)
    kicker = doc.add_paragraph(GENRE_LABEL.get(tome.genre, "Material educativo").upper())
    kicker.runs[0].font.color.rgb, kicker.runs[0].bold = accent, True
    doc.add_paragraph(tome.title, style="Title")
    doc.add_paragraph(tome.audience)
    doc.add_paragraph("Nombre: ______________________________")
    if tome.author:
        doc.add_paragraph(f"Elaborado por {tome.author}")
    if len(tome.sections) > 1:
        _page_break(doc)
        doc.add_heading("Contenido", level=1)
        for s in tome.sections:
            doc.add_paragraph(f"{s.position}. {s.title}")

    for s in tome.sections:
        _page_break(doc)
        doc.add_heading(f"{s.position}. {s.title}", level=1)
        n = 0
        for b in s.content.blocks:
            if b.type == "exercise":
                n += 1
            _block(doc, b, n, column, accent, c, drawings.get(f"{s.position}:{b.id}"), tome)

    answers = _answers(tome) if include_answer_key else []
    if answers:
        _page_break(doc)
        doc.add_heading("Solucionario", level=1)
        current = None
        for section, prompt, answer, explanation in answers:
            if section.position != current:
                current = section.position
                doc.add_heading(f"{section.position}. {section.title}", level=3)
            p = doc.add_paragraph()
            _inline(p, prompt)
            p = doc.add_paragraph()
            run = p.add_run(answer)
            run.bold = True
            if explanation:
                p.add_run(f" — {explanation}")
    if tome.sources:
        _page_break(doc)
        doc.add_heading("Fuentes", level=1)
        for src in tome.sources:
            doc.add_paragraph(f"{src.get('title', '')} — {src.get('url', '')}", style="List Bullet")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _block(doc, b, n, column, accent, c, drawing, tome) -> None:
    if b.type == "heading":
        doc.add_heading(_plain(b.text), level=b.level)
    elif b.type == "paragraph":
        _inline(doc.add_paragraph(), b.text)
    elif b.type == "list":
        style = {"number": "List Number", "bullet": "List Bullet"}.get(b.style, "List Paragraph")
        for item in b.items:
            p = doc.add_paragraph(style=style)
            if b.style == "check":
                p.add_run("☐ ")
            _inline(p, item)
    elif b.type in ("callout", "example"):
        title = b.title or (CALLOUT_LABEL.get(b.variant, "") if b.type == "callout" else "Ejemplo")
        cell = _box(doc, c["soft"])
        if title:
            run = cell.paragraphs[0].add_run(title.upper())
            run.bold, run.font.size, run.font.color.rgb = True, Pt(9), accent
            _inline(cell.add_paragraph(), b.text)
        else:
            _inline(cell.paragraphs[0], b.text)
    elif b.type == "table":
        if b.caption:
            doc.add_paragraph(b.caption).runs[0].italic = True
        table = doc.add_table(rows=1, cols=len(b.headers), style="Table Grid")
        for cell, text in zip(table.rows[0].cells, b.headers):
            _inline(cell.paragraphs[0], f"**{text}**")
        for row in b.rows:
            for cell, text in zip(table.add_row().cells, row):
                _inline(cell.paragraphs[0], text)
    elif b.type == "activity":
        cell = _box(doc, c["soft"])
        run = cell.paragraphs[0].add_run(b.title + (f"  ·  {b.duration_minutes} min" if b.duration_minutes else ""))
        run.bold = True
        if b.materials:
            cell.add_paragraph(f"Materiales: {', '.join(b.materials)}").runs[0].italic = True
        for i, step in enumerate(b.steps, 1):
            _inline(cell.add_paragraph(), f"{i}. {step}")
    elif b.type == "exercise":
        p = doc.add_paragraph()
        run = p.add_run(f"{n}. ")
        run.bold, run.font.color.rgb = True, accent
        _inline(p, f"**{b.prompt}**")
        options = b.options if b.kind == "multiple_choice" and b.options else (
            ["Verdadero", "Falso"] if b.kind == "true_false" else [])
        for o in options:
            _inline(doc.add_paragraph(), f"○ {o}")
        _lines(doc, b.answer_space_lines or (3 if b.kind == "open" else 0))
    elif b.type == "letter_grid":
        if b.instruction:
            _inline(doc.add_paragraph(), f"**{b.instruction}**")
        cells = letter_grid_cells(b)
        table = doc.add_table(rows=b.rows, cols=b.cols, style="Table Grid")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for (text, _), cell in zip(cells, (cell for row in table.rows for cell in row.cells)):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(text).font.size = Pt(22)
    elif b.type == "image":
        data = tome.images.get(b.asset_id)[0] if b.asset_id in tome.images else None
        picture = drawing or (_word_image(data) if data else None)
        if picture:
            width = {"full": 1.0, "half": 0.6, "small": 0.35}.get(b.size, 0.6)
            doc.add_picture(io.BytesIO(picture), width=int(column * width))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            doc.add_paragraph(f"[{'Dibujo para colorear' if b.kind == 'line_art' else 'Ilustración'}: {b.subject}]")
        if b.caption:
            doc.add_paragraph(b.caption).alignment = WD_ALIGN_PARAGRAPH.CENTER
        if tome.credits.get(b.asset_id):
            doc.add_paragraph(tome.credits[b.asset_id]).runs[0].font.size = Pt(8)
    elif b.type == "writing_lines":
        if b.label:
            doc.add_paragraph(b.label)
        _lines(doc, b.lines)
    elif b.type == "page_break":
        _page_break(doc)
    elif b.type in DRAWN and drawing:
        if b.type == "tracing" and b.instruction:
            _inline(doc.add_paragraph(), f"**{b.instruction}**")
        doc.add_picture(io.BytesIO(drawing), width=column)
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if getattr(b, "caption", None):
            doc.add_paragraph(b.caption).alignment = WD_ALIGN_PARAGRAPH.CENTER


def _rasterise(tome: Tome, c: dict, head: str) -> dict[str, bytes]:
    """"<position>:<block id>" -> PNG of every code-drawn block and every vector colouring page."""
    svgs: dict[str, str] = {}
    for s in tome.sections:
        for b in s.content.blocks:
            key = f"{s.position}:{b.id}"
            if b.type in DIAGRAM_TYPES:
                svgs[key] = str(render_diagram(b, c, head))
            elif b.type in WORKSHEET_TYPES:
                svgs[key] = str(render_worksheet(b, c, head, tome.image_uri))
            elif b.type == "keyboard":
                svgs[key] = str(keyboard_svg(b, c["tint"], c["accent"]))
            elif b.type == "tracing":
                svgs[key] = str(tracing_svg(b))
            elif b.type == "image" and b.asset_id in tome.images and tome.images[b.asset_id][1] == "image/svg+xml":
                svgs[key] = tome.images[b.asset_id][0].decode()
    if not svgs:
        return {}
    from weasyprint import HTML
    from weasyprint.text.fonts import FontConfiguration

    pages, sizes = [], []
    for i, svg in enumerate(svgs.values()):
        m = VIEWBOX.search(svg)
        vw, vh = (float(m.group(3)), float(m.group(4))) if m else (640.0, 480.0)
        sizes.append((vw, vh))
        pages.append(f'<style>@page p{i} {{ size: {vw}px {vh}px; margin: 0 }}</style>'
                     f'<div style="page: p{i}; width: {vw}px; height: {vh}px; break-after: page">'
                     f'{_sized(svg, vw, vh)}</div>')
    fonts = (ASSETS / "fonts").as_uri()
    faces = "".join(f'@font-face {{ font-family: "{name}"; src: url("{fonts}/{file}"); }}'
                    for name, (file, _) in FONTS.items())
    faces += (f'@font-face {{ font-family: "Andika"; src: url("{fonts}/Andika-Regular.ttf"); }}'
              f'@font-face {{ font-family: "Andika"; font-weight: 700; src: url("{fonts}/Andika-Bold.ttf"); }}'
              f'@font-face {{ font-family: "Noto Music"; src: url("{fonts}/NotoMusic-Regular.ttf"); }}')
    html = f"<html><head><style>{faces} body {{ margin: 0 }} svg {{ display: block }}</style></head><body>" \
           f"{''.join(pages)}</body></html>"
    config = FontConfiguration()
    pdf = pypdfium2.PdfDocument(HTML(string=html).write_pdf(font_config=config))
    out = {}
    for key, page, (vw, _) in zip(svgs, pdf, sizes):
        image = page.render(scale=PNG_WIDTH / (vw * 0.75), fill_color=(255, 255, 255, 255)).to_pil()
        buf = io.BytesIO()
        image.convert("RGB").save(buf, "PNG", optimize=True)
        out[key] = buf.getvalue()
    return out


def _sized(svg: str, w: float, h: float) -> str:
    """The drawing filling its page exactly, whatever width/height it was drawn with."""
    svg = re.sub(r'\s(width|height)="[^"]*"', "", svg[:svg.index(">") + 1], count=2) + svg[svg.index(">") + 1:]
    return svg.replace("<svg", f'<svg width="{w}px" height="{h}px"', 1)


def _word_image(data: bytes) -> bytes:
    """Word reads PNG and JPEG; WebP and others are converted."""
    if data[:3] == b"\xff\xd8\xff" or data[:8] == b"\x89PNG\r\n\x1a\n":
        return data
    buf = io.BytesIO()
    Image.open(io.BytesIO(data)).convert("RGB").save(buf, "PNG")
    return buf.getvalue()


def _answers(tome: Tome) -> list[tuple]:
    out = []
    for s in tome.sections:
        for b in s.content.blocks:
            if b.type == "exercise" and b.answer:
                out.append((s, _plain(b.prompt), b.answer, b.explanation))
            for prompt, answer in worksheet_answers(b):
                out.append((s, prompt, answer, None))
    return out


def _inline(paragraph, text: str) -> None:
    """**bold** and *italic*, the only markup the Writer may use."""
    for part in re.split(r"(\*\*.+?\*\*|(?<!\*)\*(?!\s).+?(?<!\s)\*(?!\*))", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _plain(text: str) -> str:
    return re.sub(r"\*{1,2}(.+?)\*{1,2}", r"\1", text)


def _box(doc, fill: str):
    """A one-cell shaded table: Word's closest thing to the PDF's callout boxes."""
    cell = doc.add_table(rows=1, cols=1).rows[0].cells[0]
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shading)
    doc.add_paragraph()
    return cell


def _lines(doc, count: int) -> None:
    for _ in range(count):
        doc.add_paragraph("_" * 70)


def _page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
