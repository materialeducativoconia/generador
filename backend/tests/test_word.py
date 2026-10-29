"""Word export: editable text for what the Writer wrote, pictures for what the code draws."""

import io

from docx import Document

from app.word import render_docx
from tests.test_render import tome

BLOCKS = [
    {"id": "b1", "type": "paragraph", "text": "Hoy **sumamos** con *dibujos*."},
    {"id": "b2", "type": "callout", "variant": "tip", "text": "Usa los dedos."},
    {"id": "b3", "type": "exercise", "kind": "multiple_choice", "prompt": "¿Cuánto es 1 + 3?",
     "options": ["3", "4"], "answer": "4"},
    {"id": "b4", "type": "tracing", "text": "Mi gato se llama Toby.", "repeat": 2},
    {"id": "b5", "type": "number_line", "start": 0, "end": 5, "blanks": [1, 3]},
    {"id": "b6", "type": "picture_math", "operations": [{"object": "manzana", "a": 2, "op": "+", "b": 1}]},
    {"id": "b7", "type": "table", "headers": ["Número", "Dedos"], "rows": [["2", "dos"]]},
]


def test_word_keeps_text_editable_and_draws_the_rest():
    data = render_docx(tome(*BLOCKS, genre="worksheets"), paper="A4", include_answer_key=True)
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)

    assert "Hoy sumamos con dibujos." in text and "○ 4" in text
    assert any(r.bold and r.text == "sumamos" for p in doc.paragraphs for r in p.runs)
    assert "Solucionario" in text and "2 + 1" in text  # the computed answer of the picture maths
    assert len(doc.inline_shapes) == 3  # tracing, number line, picture maths
    assert doc.tables[-1].rows[1].cells[1].text == "dos"
    assert round(doc.sections[0].page_width.mm) == 210


def test_word_without_answer_key():
    data = render_docx(tome(*BLOCKS[:3]), paper="Letter", include_answer_key=False)

    assert "Solucionario" not in "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
